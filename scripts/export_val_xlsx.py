"""Export a search-validation run (data/eval/results/val_<stamp>.json) to an Excel workbook.

  python scripts/export_val_xlsx.py                       # latest run
  python scripts/export_val_xlsx.py data/eval/results/val_20261006-1604.json
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "data" / "eval" / "results"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
SUB_FILL = PatternFill("solid", start_color="DDEBF7")
GOOD = PatternFill("solid", start_color="C6EFCE")
MID = PatternFill("solid", start_color="FFEB9C")
BAD = PatternFill("solid", start_color="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VARIANTS = [("en", "Tiếng Anh"), ("vi", "Tiếng Việt có dấu"), ("vi_noacc", "Không dấu"), ("vi_chat", "Kiểu chat")]
CATEGORY_VI = {"product": "Sản phẩm", "insight": "Insight", "compliance": "Compliance", "sop": "SOP",
               "policy": "Chính sách", "creative": "Creative", "brand": "Brand", "gap": "Khoảng trống dữ liệu"}


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, first_row, last_row, ncols, wrap_cols=()):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap_cols)


def widths(ws, values):
    for i, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def traffic_light(ws, rng, good, bad):
    """Green >= good, red < bad, yellow in between (for 0-1 scores)."""
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[str(good)], fill=GOOD))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(bad)], fill=BAD))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=[str(bad), str(good)], fill=MID))


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else sorted(RESULTS.glob("val_*.json"))[-1]
    data = json.loads(path.read_text(encoding="utf-8"))
    questions = {json.loads(l)["id"]: json.loads(l)
                 for l in (ROOT / "data" / "eval" / "val_kb.jsonl").read_text(encoding="utf-8").splitlines() if l}
    reviews_q = {json.loads(l)["id"]: json.loads(l)
                 for l in (ROOT / "data" / "eval" / "val_reviews.jsonl").read_text(encoding="utf-8").splitlines() if l}
    stamp = path.stem.removeprefix("val_")

    wb = Workbook()

    # ---------- Sheet 2: KB chi tiết (long format, one row per query run) ----------
    det = wb.create_sheet("KB - chi tiết")
    title(det, "Tìm tài liệu trong kb: từng lượt chạy",
          "Hạng = vị trí của chunk đúng đầu tiên trong top 10 (trống = không tìm thấy). Top 5 = 1 nếu hạng ≤ 5.")
    head = ["ID", "Nhóm", "Biến thể", "Câu hỏi", "Hạng", "Top 1", "Top 5", "Đáp án mong đợi",
            "Kết quả 1", "Kết quả 2", "Kết quả 3"]
    det.append([])
    det.append(head)
    hr = 4
    style_header(det, hr, len(head))
    vname = dict(VARIANTS)
    for r in data["kb"]:
        det.append([r["id"], CATEGORY_VI[r["category"]], vname[r["variant"]], r["query"], r["rank"], None, None,
                    "\n".join(questions[r["id"]]["expected"]), *(r["top3"] + ["", "", ""])[:3]])
        row = det.max_row
        det.cell(row=row, column=6, value=f'=IF(E{row}=1,1,0)')
        det.cell(row=row, column=7, value=f'=IF(AND(E{row}<>"",E{row}<=5),1,0)')
    det_first, det_last = hr + 1, det.max_row
    style_body(det, det_first, det_last, len(head), wrap_cols=(4, 8, 9, 10, 11))
    for row in range(det_first, det_last + 1):
        for c in (5, 6, 7):
            det.cell(row=row, column=c).alignment = Alignment(horizontal="center", vertical="top")
    det.conditional_formatting.add(f"G{det_first}:G{det_last}", CellIsRule(operator="equal", formula=["1"], fill=GOOD))
    det.conditional_formatting.add(f"G{det_first}:G{det_last}", CellIsRule(operator="equal", formula=["0"], fill=BAD))
    widths(det, [7, 14, 16, 48, 7, 7, 7, 40, 40, 40, 40])
    det.freeze_panes = det.cell(row=hr + 1, column=5)
    det.auto_filter.ref = f"A{hr}:{get_column_letter(len(head))}{det_last}"
    D = "'KB - chi tiết'"

    # ---------- Sheet 3: KB theo câu (wide: one row per question, rank per variant) ----------
    wide = wb.create_sheet("KB - theo câu")
    title(wide, "Tìm tài liệu trong kb: hạng của chunk đúng theo từng câu",
          "Xanh = hạng 1, vàng = hạng 2-5, đỏ = ngoài top 5 hoặc không tìm thấy (ô trống).")
    head = ["ID", "Nhóm", "Câu hỏi (tiếng Việt)"] + [v for _, v in VARIANTS] + ["Đáp án mong đợi"]
    wide.append([])
    wide.append(head)
    style_header(wide, 4, len(head))
    ranks = {(r["id"], r["variant"]): r["rank"] for r in data["kb"]}
    for qid, q in questions.items():
        wide.append([qid, CATEGORY_VI[q["category"]], q["vi"]] + [ranks.get((qid, v)) for v, _ in VARIANTS]
                    + ["\n".join(q["expected"])])
    w_last = wide.max_row
    style_body(wide, 5, w_last, len(head), wrap_cols=(3, 8))
    rank_rng = f"D5:G{w_last}"
    wide.conditional_formatting.add(rank_rng, CellIsRule(operator="equal", formula=["1"], fill=GOOD))
    wide.conditional_formatting.add(rank_rng, CellIsRule(operator="between", formula=["2", "5"], fill=MID))
    wide.conditional_formatting.add(rank_rng, CellIsRule(operator="greaterThan", formula=["5"], fill=BAD))
    from openpyxl.formatting.rule import FormulaRule
    wide.conditional_formatting.add(rank_rng, FormulaRule(formula=["ISBLANK(D5)"], fill=BAD))
    for row in range(5, w_last + 1):
        for c in range(4, 8):
            wide.cell(row=row, column=c).alignment = Alignment(horizontal="center", vertical="top")
    widths(wide, [7, 18, 55, 12, 12, 12, 12, 45])
    wide.freeze_panes = "D5"
    wide.auto_filter.ref = f"A4:{get_column_letter(len(head))}{w_last}"

    # ---------- Sheet 4: Review ----------
    rv = wb.create_sheet("Review")
    title(rv, "Tìm review có bộ lọc: precision@5",
          "Precision@5 = số review trong top 5 khớp từ khoá nhãn / 5. Vi phạm = số kết quả sai bộ lọc (phải bằng 0).")
    head = ["ID", "Biến thể", "Câu hỏi", "Bộ lọc", "Từ khoá nhãn (regex)", "Số kết quả", "Precision@5",
            "Vi phạm bộ lọc", "Review không liên quan"]
    rv.append([])
    rv.append(head)
    style_header(rv, 4, len(head))
    rvname = {"en": "Tiếng Anh", "vi": "Tiếng Việt có dấu", "vi_noacc": "Không dấu"}
    for r in data["reviews"]:
        q = reviews_q[r["id"]]
        rv.append([r["id"], rvname[r["variant"]], r["query"], json.dumps(q["filter"], ensure_ascii=False),
                   q["relevant"], r["n_hits"], r["precision@5"], r["violations"], "\n".join(r["irrelevant"])])
    rv_last = rv.max_row
    style_body(rv, 5, rv_last, len(head), wrap_cols=(3, 4, 5, 9))
    for row in range(5, rv_last + 1):
        rv.cell(row=row, column=7).number_format = "0%"
        for c in (6, 7, 8):
            rv.cell(row=row, column=c).alignment = Alignment(horizontal="center", vertical="top")
    traffic_light(rv, f"G5:G{rv_last}", 0.8, 0.5)
    rv.conditional_formatting.add(f"H5:H{rv_last}", CellIsRule(operator="greaterThan", formula=["0"], fill=BAD))
    widths(rv, [7, 16, 38, 34, 30, 9, 11, 10, 30])
    rv.freeze_panes = "C5"
    rv.auto_filter.ref = f"A4:{get_column_letter(len(head))}{rv_last}"
    R = "Review"

    # ---------- Sheet 5: Ngoài phạm vi ----------
    sc = wb.create_sheet("Ngoài phạm vi")
    title(sc, "Nhận biết câu hỏi ngoài phạm vi bằng điểm tương đồng",
          "Đổi ngưỡng ở ô C4 để xem bao nhiêu câu ngoài phạm vi bị loại (điểm < ngưỡng).")
    sc["A4"] = "Ngưỡng điểm tương đồng (cosine)"
    sc["A4"].font = Font(name=FONT, bold=True)
    sc["C4"] = round(data["scope_summary"]["threshold"], 3)
    sc["C4"].font = Font(name=FONT, color="0000FF", bold=True)
    sc["C4"].fill = PatternFill("solid", start_color="FFFF00")
    sc["D4"] = "Ngưỡng tốt nhất tìm được trên lần chạy này (giữ 100% câu hợp lệ)."
    sc["D4"].font = Font(name=FONT, italic=True, size=9, color="595959")
    sc["A5"] = "Tỉ lệ câu ngoài phạm vi bị loại"
    sc["A5"].font = Font(name=FONT, bold=True)
    head = ["ID", "Ngôn ngữ", "Lý do ngoài phạm vi", "Câu hỏi", "Điểm (cosine)", "Bị loại?", "Chunk gần nhất"]
    sc.append([])
    sc.append(head)
    style_header(sc, 7, len(head))
    for r in sorted(data["scope"], key=lambda r: -r["cosine"]):
        sc.append([r["id"], "Tiếng Anh" if r["variant"] == "en" else "Tiếng Việt", r["reason"], r["query"],
                   round(r["cosine"], 3), None, r["top"]])
        row = sc.max_row
        sc.cell(row=row, column=6, value=f'=IF(E{row}<$C$4,"Có","Không")')
    sc_last = sc.max_row
    sc["C5"] = f'=COUNTIF(F8:F{sc_last},"Có")/COUNTA(F8:F{sc_last})'
    sc["C5"].number_format = "0%"
    sc["C5"].font = Font(name=FONT, bold=True)
    style_body(sc, 8, sc_last, len(head), wrap_cols=(3, 4, 7))
    for row in range(8, sc_last + 1):
        sc.cell(row=row, column=5).number_format = "0.000"
        sc.cell(row=row, column=5).alignment = Alignment(horizontal="center", vertical="top")
        sc.cell(row=row, column=6).alignment = Alignment(horizontal="center", vertical="top")
    sc.conditional_formatting.add(f"F8:F{sc_last}", CellIsRule(operator="equal", formula=['"Có"'], fill=GOOD))
    sc.conditional_formatting.add(f"F8:F{sc_last}", CellIsRule(operator="equal", formula=['"Không"'], fill=BAD))
    s = data["scope_summary"]
    ref_row = sc_last + 2
    sc.cell(row=ref_row, column=1, value="Tham chiếu: điểm của câu hỏi hợp lệ (phần A, tiếng Anh + tiếng Việt)").font = Font(name=FONT, bold=True)
    for i, (k, lab) in enumerate((("min", "Thấp nhất"), ("median", "Trung vị"), ("max", "Cao nhất"))):
        sc.cell(row=ref_row + 1 + i, column=1, value=lab).font = Font(name=FONT)
        c = sc.cell(row=ref_row + 1 + i, column=3, value=round(s["in_scope"][k], 3))
        c.font = Font(name=FONT, color="0000FF")
        c.number_format = "0.000"
    widths(sc, [7, 12, 26, 46, 13, 10, 48])
    sc.freeze_panes = "A8"

    # ---------- Sheet 1: Tổng quan (formulas over the detail sheets) ----------
    ov = wb.active
    ov.title = "Tổng quan"
    title(ov, f"Kết quả validation search ({stamp[:8]} {stamp[9:11]}:{stamp[11:]})",
          "Model e5-large đa ngôn ngữ + BM25 + từ điển Việt-Anh. Bộ held-out, không dùng để chọn model hay viết từ điển.")
    ov["A4"] = "A. Tìm tài liệu trong kb (53 câu)"
    ov["A4"].font = Font(name=FONT, bold=True, size=12)
    head = ["Biến thể", "Số lượt", "Hit@1", "Hit@5"]
    for c, h in enumerate(head, 1):
        ov.cell(row=5, column=c, value=h)
    style_header(ov, 5, len(head))
    rng = lambda col: f"{D}!${col}${det_first}:${col}${det_last}"
    r0 = 6
    for i, (_, lab) in enumerate(VARIANTS + [("all", "Tất cả")]):
        row = r0 + i
        ov.cell(row=row, column=1, value=lab)
        if lab == "Tất cả":
            ov.cell(row=row, column=2, value=f"=COUNTA({rng('C')})")
            ov.cell(row=row, column=3, value=f"=AVERAGE({rng('F')})")
            ov.cell(row=row, column=4, value=f"=AVERAGE({rng('G')})")
        else:
            ov.cell(row=row, column=2, value=f'=COUNTIF({rng("C")},A{row})')
            ov.cell(row=row, column=3, value=f'=AVERAGEIF({rng("C")},A{row},{rng("F")})')
            ov.cell(row=row, column=4, value=f'=AVERAGEIF({rng("C")},A{row},{rng("G")})')
    a_last = r0 + len(VARIANTS)
    style_body(ov, r0, a_last, len(head))
    for row in range(r0, a_last + 1):
        for c in (3, 4):
            ov.cell(row=row, column=c).number_format = "0%"
    for c in range(1, 5):
        ov.cell(row=a_last, column=c).font = Font(name=FONT, bold=True, size=10)
    traffic_light(ov, f"C{r0}:D{a_last}", 0.85, 0.7)

    # hit@5 by category x variant
    cat_row = a_last + 2
    ov.cell(row=cat_row, column=1, value="Hit@5 theo nhóm câu hỏi").font = Font(name=FONT, bold=True, size=12)
    head = ["Nhóm"] + [v for _, v in VARIANTS]
    for c, h in enumerate(head, 1):
        ov.cell(row=cat_row + 1, column=c, value=h)
    style_header(ov, cat_row + 1, len(head))
    for i, cat in enumerate(CATEGORY_VI.values()):
        row = cat_row + 2 + i
        ov.cell(row=row, column=1, value=cat)
        for j, (_, lab) in enumerate(VARIANTS):
            col = get_column_letter(2 + j)
            ov.cell(row=row, column=2 + j,
                    value=f'=AVERAGEIFS({rng("G")},{rng("B")},$A{row},{rng("C")},{col}${cat_row + 1})')
            ov.cell(row=row, column=2 + j).number_format = "0%"
    c_last = cat_row + 1 + len(CATEGORY_VI)
    style_body(ov, cat_row + 2, c_last, len(head))
    traffic_light(ov, f"B{cat_row + 2}:E{c_last}", 0.85, 0.7)

    # reviews summary
    b_row = c_last + 2
    ov.cell(row=b_row, column=1, value="B. Tìm review có bộ lọc (15 câu)").font = Font(name=FONT, bold=True, size=12)
    head = ["Biến thể", "Số câu", "Precision@5", "Vi phạm bộ lọc"]
    for c, h in enumerate(head, 1):
        ov.cell(row=b_row + 1, column=c, value=h)
    style_header(ov, b_row + 1, len(head))
    rr = lambda col: f"{R}!${col}$5:${col}${rv_last}"
    for i, lab in enumerate(rvname.values()):
        row = b_row + 2 + i
        ov.cell(row=row, column=1, value=lab)
        ov.cell(row=row, column=2, value=f'=COUNTIF({rr("B")},A{row})')
        ov.cell(row=row, column=3, value=f'=AVERAGEIF({rr("B")},A{row},{rr("G")})')
        ov.cell(row=row, column=3).number_format = "0%"
        ov.cell(row=row, column=4, value=f'=SUMIF({rr("B")},A{row},{rr("H")})')
    b_last = b_row + 1 + len(rvname)
    style_body(ov, b_row + 2, b_last, len(head))
    traffic_light(ov, f"C{b_row + 2}:C{b_last}", 0.85, 0.7)

    # scope summary
    s_row = b_last + 2
    ov.cell(row=s_row, column=1, value="C. Câu hỏi ngoài phạm vi (12 câu x 2 ngôn ngữ)").font = Font(name=FONT, bold=True, size=12)
    rows = [("Ngưỡng cosine đang dùng", "='Ngoài phạm vi'!C4", "0.000"),
            ("Tỉ lệ câu ngoài phạm vi bị loại", "='Ngoài phạm vi'!C5", "0%"),
            ("Điểm cao nhất của câu ngoài phạm vi", f"=MAX('Ngoài phạm vi'!E8:E{sc_last})", "0.000"),
            ("Trung vị điểm câu hợp lệ", f"='Ngoài phạm vi'!C{ref_row + 2}", "0.000")]
    for i, (lab, f, fmt) in enumerate(rows):
        row = s_row + 1 + i
        ov.cell(row=row, column=1, value=lab)
        c = ov.cell(row=row, column=3, value=f)
        c.number_format = fmt
        c.font = Font(name=FONT, color="008000")
    style_body(ov, s_row + 1, s_row + len(rows), 4)
    for i in range(len(rows)):
        ov.cell(row=s_row + 1 + i, column=3).font = Font(name=FONT, size=10, color="008000")
    note = s_row + len(rows) + 2
    ov.cell(row=note, column=1, value="Màu: xanh ≥ 85%, vàng 70-85%, đỏ < 70%. Chữ xanh lá = liên kết sang sheet khác.")
    ov.cell(row=note, column=1).font = Font(name=FONT, italic=True, size=9, color="595959")
    widths(ov, [34, 16, 16, 16, 16])
    ov.sheet_view.showGridLines = False

    # No cached values are written; make Excel compute every formula when the file is opened.
    wb.calculation.fullCalcOnLoad = True
    out = RESULTS / f"val_{stamp}.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
